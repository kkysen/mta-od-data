"""Scale the OD data's weekdays by each station's change in entries since another year.

The OD data starts in 2023, but the MTA's annual station ridership spreadsheet
has each station complex's counted average weekday entries back to 2019.
Each complex grows by its counted entries in the target year
over those in the OD data's year, both from the spreadsheet,
so the OD data only supplies the pattern of who goes where, not a level to compare.
Both count swipes and taps, and so does the OD data
(scaled to each origin's swipes, not for fare evasion),
so none of them include fare evaders.
Scaling an OD slice by iterative proportional fitting (Fratar)
to its own trips from and to each complex, grown,
keeps the slice's pattern, closures, and seasonality.
A station's daily exits are taken to grow like its entries,
the turnstile exit counts being incomplete (exits through emergency gates).
"""

import csv
from collections import defaultdict
from collections.abc import Iterable
from pathlib import Path
from typing import Annotated
from urllib.request import urlretrieve

from openpyxl import load_workbook
from typer import Option, Typer

from mta_od_data import DATA
from mta_od_data.analyze.common import connection
from mta_od_data.prepare import DEFAULT_PARQUET, RIDERSHIP_DECIMAL

app = Typer()

# The 2025 subway ridership data, from
# https://www.mta.info/agency/new-york-city-transit/ridership/2025:
# average weekday, weekend, and annual entries per station complex, 2019 to 2025.
SPREADSHEET_URL = "https://www.mta.info/document/213556"
DEFAULT_SPREADSHEET = DATA / "mta_subway_ridership_2019_2025.xlsx"
# The spreadsheet's station names, matched to complex IDs by hand.
# The spreadsheet's Times Sq row also counts 42 St-Bryant Pk/5 Av,
# a separate complex here, so that row maps to both, both growing as the row does.
MAPPING = Path(__file__).parent / "station_ridership_complexes.csv"
# Iterating until every complex's trips from and to it are within this of its target,
# or for at most so many iterations, when a slice can't fit exactly.
TOLERANCE = 1e-6
MAX_ITERATIONS = 30


def average_weekday_entries(spreadsheet: Path, year: int) -> dict[str, float]:
    """Each station row's average weekday entries in `year`, by its name."""
    workbook = load_workbook(spreadsheet, read_only=True)
    rows = list(workbook["Avg Weekday"].iter_rows(values_only=True))
    header = rows[1]
    column = header.index(year)
    entries: dict[str, float] = {}
    for row in rows[2:]:
        name, borough, value = row[0], row[2], row[column]
        # Borough subtotals and the systemwide total have no borough.
        if isinstance(name, str) and borough and isinstance(value, int | float):
            entries[name] = float(value)
    return entries


def station_complexes() -> dict[str, list[int]]:
    """Each spreadsheet station row's complexes."""
    complexes_of: dict[str, list[int]] = defaultdict(list)
    with MAPPING.open() as f:
        for row in csv.DictReader(f):
            complexes_of[row["station"]].append(int(row["complex_id"]))
    return complexes_of


def complex_growth(
    to: dict[str, float], since: dict[str, float]
) -> tuple[dict[int, float], float]:
    """Each complex's growth, `to` over `since` entries of its station row,
    and the systemwide growth, for complexes the spreadsheet doesn't have."""
    complexes_of = station_complexes()
    if missing := sorted(set(to) - set(complexes_of)):
        raise ValueError(f"{MAPPING}: no complex for {missing}")
    growth = {c: to[s] / since[s] for s in to for c in complexes_of[s]}
    return growth, sum(to.values()) / sum(since.values())


def od_entries(rows: Iterable[str]) -> dict[str, float]:
    """Each station row's average weekday trips from it in the `daily` table:
    the OD data's own count of its entries, averaged over the slices."""
    complexes_of = station_complexes()
    trips_from = dict(
        connection()
        .execute(
            """
            SELECT o, sum(v) / (SELECT count(DISTINCT (month, dow)) FROM daily)
            FROM daily GROUP BY o
            """
        )
        .fetchall()
    )
    return {s: sum(trips_from.get(c, 0.0) for c in complexes_of[s]) for s in rows}


@app.command()
def scale_od(
    year: Annotated[int, Option(help="The year whose station entries to scale to")],
    out: Annotated[Path, Option(help="Output Parquet, in the OD Parquet's schema")],
    od: Annotated[Path, Option(help="The OD Parquet")] = DEFAULT_PARQUET,
    od_year: Annotated[int, Option(help="The OD data's year to scale")] = 2025,
    spreadsheet: Annotated[
        Path, Option(help="The MTA station ridership spreadsheet (fetched if missing)")
    ] = DEFAULT_SPREADSHEET,
    to_counts: Annotated[
        bool,
        Option(
            help=(
                "Grow each complex to its counted `--year` entries "
                "from the OD data's own `--od-year` level, "
                "not by its counted growth since `--od-year`: "
                "the OD data runs about 1% under the counts, more in Manhattan"
            )
        ),
    ] = False,
) -> None:
    """Scale `--od-year`'s weekday OD slices to `--year`'s station entries.

    Each complex grows by its counted `--year` over `--od-year` entries,
    applied to each (month, weekday) slice's own trips from and to it,
    so a slice's own closures and seasonality stay.

    With `--to-counts`, each complex grows to its counted `--year` entries
    from the OD data's own `--od-year` level instead,
    so `--year` can be `--od-year` itself, to level the OD data with the counts.

    \b
    Examples:
        mta-od-data scale-od --year 2019 --out data/mta_od_scaled_to_2019.parquet
        mta-od-data scale-od --year 2025 --to-counts \
            --out data/mta_od_2025_counted.parquet
    """
    if not spreadsheet.exists():
        print(f"fetching {SPREADSHEET_URL}")
        urlretrieve(SPREADSHEET_URL, spreadsheet)
    con = connection()
    con.execute(
        """
        CREATE OR REPLACE TEMP TABLE daily AS
        SELECT Month AS month, "Day of Week" AS dow,
               "Origin Station Complex ID" AS o, "Destination Station Complex ID" AS d,
               sum("Estimated Average Ridership")::DOUBLE AS v
        FROM read_parquet($od)
        WHERE Year = $od_year AND "Day of Week" NOT IN ('Saturday', 'Sunday')
        GROUP BY ALL
        """,
        {"od": str(od), "od_year": od_year},
    )
    to = average_weekday_entries(spreadsheet, year)
    since = (
        od_entries(to) if to_counts else average_weekday_entries(spreadsheet, od_year)
    )
    growth, systemwide = complex_growth(to, since)
    # The subway spreadsheet has no Staten Island Railway,
    # but the OD data has its two fare-controlled stations, St George and Tompkinsville:
    # those grow systemwide.
    complexes = [c for (c,) in con.execute("SELECT DISTINCT o FROM daily").fetchall()]
    if missing := sorted(set(complexes) - set(growth)):
        print(
            f"no entries for complexes {missing}: "
            f"growing by the systemwide {systemwide:.3f}"
        )
        for c in missing:
            growth[c] = systemwide
    con.execute("CREATE OR REPLACE TEMP TABLE growth (c BIGINT, g DOUBLE)")
    con.executemany("INSERT INTO growth VALUES (?, ?)", list(growth.items()))
    # Each slice's trips from and to a complex, grown:
    # so a complex closed or busier in one slice stays so, scaled.
    # Trips to are normalized to the same total as trips from, for a feasible fit.
    con.execute(
        """
        CREATE OR REPLACE TEMP TABLE row_target AS
        SELECT month, dow, o AS c, sum(v) * g AS t
        FROM daily JOIN growth ON growth.c = o GROUP BY month, dow, o, g;
        CREATE OR REPLACE TEMP TABLE col_target AS
        WITH raw AS (
            SELECT month, dow, d AS c, sum(v) * g AS t
            FROM daily JOIN growth ON growth.c = d GROUP BY month, dow, d, g
        )
        SELECT month, dow, c,
               t * (
                   SELECT sum(t) FROM row_target r
                   WHERE r.month = raw.month AND r.dow = raw.dow
               ) / sum(t) OVER (PARTITION BY month, dow) AS t
        FROM raw;
        """
    )

    # Row factors `a` and column factors `b` per slice, alternately fitted.
    con.execute(
        """
        CREATE OR REPLACE TEMP TABLE a AS
        SELECT DISTINCT month, dow, o, 1.0::DOUBLE AS f FROM daily;
        CREATE OR REPLACE TEMP TABLE b AS
        SELECT DISTINCT month, dow, d, 1.0::DOUBLE AS f FROM daily;
        """
    )
    # Columns fit exactly after `b`; how far rows are off, and by how many riders.
    row_errors = """
        SELECT s.month, s.dow, s.o, s.s, t.t, abs(s.s / t.t - 1) AS error
        FROM (
            SELECT month, dow, o, sum(v * a.f * b.f) AS s
            FROM daily JOIN a USING (month, dow, o) JOIN b USING (month, dow, d)
            GROUP BY ALL
        ) s JOIN row_target t ON t.month = s.month AND t.dow = s.dow AND t.c = s.o
    """
    iterations = 0
    while iterations < MAX_ITERATIONS:
        iterations += 1
        con.execute(
            """
            CREATE OR REPLACE TEMP TABLE a AS
            SELECT month, dow, o, t.t / sum(v * b.f) AS f
            FROM daily JOIN b USING (month, dow, d)
            JOIN row_target t ON t.month = daily.month AND t.dow = daily.dow AND t.c = o
            GROUP BY month, dow, o, t.t;
            CREATE OR REPLACE TEMP TABLE b AS
            SELECT month, dow, d, t.t / sum(v * a.f) AS f
            FROM daily JOIN a USING (month, dow, o)
            JOIN col_target t ON t.month = daily.month AND t.dow = daily.dow AND t.c = d
            GROUP BY month, dow, d, t.t;
            """
        )
        result: tuple[float] | None = con.execute(
            f"SELECT max(error) FROM ({row_errors})"
        ).fetchone()
        assert result is not None, "aggregate query always returns one row"
        if result[0] < TOLERANCE:
            break
    # Not always exact: a complex all but cut off in a slice
    # (e.g. by a service suspension) can't always reach its target through what's left.
    summary: tuple[float, float] | None = con.execute(
        f"SELECT sum(abs(s - t)) / sum(t), max(error) FROM ({row_errors})"
    ).fetchone()
    assert summary is not None, "aggregate query always returns one row"
    print(
        f"after {iterations} iterations: trips from complexes off target by "
        f"{summary[0]:.2e} of all trips, at most {summary[1]:.2e} of one complex's"
    )
    if summary[1] >= TOLERANCE:
        worst = con.execute(
            f"SELECT * FROM ({row_errors}) ORDER BY error DESC LIMIT 5"
        ).fetchall()
        print(f"worst (month, day, complex, trips from, target, error): {worst}")

    con.execute(
        f"""
        COPY (
            SELECT p.Year, p.Month, p."Day of Week", p."Hour of Day",
                   p."Origin Station Complex ID", p."Destination Station Complex ID",
                   (p."Estimated Average Ridership" * a.f * b.f)::{RIDERSHIP_DECIMAL}
                       AS "Estimated Average Ridership"
            FROM read_parquet($od) p
            JOIN a ON a.month = p.Month AND a.dow = p."Day of Week"
                  AND a.o = p."Origin Station Complex ID"
            JOIN b ON b.month = p.Month AND b.dow = p."Day of Week"
                  AND b.d = p."Destination Station Complex ID"
            WHERE p.Year = $od_year
            ORDER BY "Origin Station Complex ID", "Destination Station Complex ID",
                     "Day of Week", "Hour of Day"
        ) TO $out (FORMAT parquet, COMPRESSION zstd, PARQUET_VERSION v2)
        """,
        {"od": str(od), "od_year": od_year, "out": str(out)},
    )
    total: tuple[float] | None = con.execute(
        """
        SELECT avg(s) FROM (
            SELECT Month, "Day of Week", sum("Estimated Average Ridership") s
            FROM read_parquet($out) GROUP BY ALL
        )
        """,
        {"out": str(out)},
    ).fetchone()
    assert total is not None, "aggregate query always returns one row"
    print(
        f"wrote {out}: {total[0]:,.0f} riders a weekday slice on average "
        f"(systemwide growth {systemwide:.3f})"
    )
